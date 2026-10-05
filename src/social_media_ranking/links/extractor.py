"""Read the creator Excel sheet and scrape each bio-link page (section 3, step A)."""
from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Callable, Dict, Iterable, List, Optional, Sequence, Tuple
from urllib.parse import urljoin

from ..models import CONFIRMATION_DEFAULT, STATUS_INVALID, STATUS_RESOLVED, CreatorProfile
from ..storage import DataRepo, today_stamp, write_csv
from .validator import IGNORED, LinkCandidate, ValidationResult, validate_candidates

log = logging.getLogger(__name__)

# Column headers accepted for the creator name and the bio-link URL.
NAME_COLUMNS = ("creator", "creator name", "name", "influencer", "full name", "creator_name")
LINK_COLUMNS = ("linktree", "linktree url", "link", "url", "bio link", "biolink", "links", "profile")
# Optional columns (section 3) — absent means every creator defaults to Pending / "".
CONFIRMATION_COLUMNS = ("confirmation", "status", "confirmed")
LANGUAGE_COLUMNS = ("language", "lang")
# Contact info (section 2: "never guess or invent", and these never leave
# this repo — export_public.py's allowlist doesn't reference them, and
# tests/test_export_public.py's contact scan guards against a regression).
# "emali"/"whastapp" are real typos found in the live roster sheet, not
# hypothetical — a substring match alone misses both ("emali" doesn't
# contain "email"; "whastapp" doesn't contain "whatsapp").
EMAIL_COLUMNS = ("email", "e-mail", "email address", "emali")
WHATSAPP_COLUMNS = ("whatsapp", "phone", "whatsapp number", "phone number", "whastapp")

_URL_IN_TEXT = re.compile(r"https?://[^\s\"'<>)\]]+")

# Mistakes confirmed in the live roster sheet (not hypothetical): someone
# pasted the browser tab's title instead of the URL, "lintr.ee" as a typo
# for "linktr.ee", and a bare "linktr.ee/..." with no scheme at all (which
# fails outright at fetch time -- requests has no base URL to resolve
# against, unlike the outbound links validator.normalize_url() handles
# later for links *found on* a page). Fixed narrowly, not Desktop/ranking's
# stricter "discard anything that isn't linktr.ee" rule -- this pipeline's
# link scraping isn't Linktree-specific (only the avatar/language pull is),
# so a legitimate non-Linktree bio-link page shouldn't be dropped here.
_LINKTREE_TITLE_PASTE = re.compile(r"^([A-Za-z0-9._-]+)\s+Official:", re.IGNORECASE)
_LINTR_TYPO = re.compile(r"^https?://(?:www\.)?lintr\.ee/", re.IGNORECASE)


def _clean_biolink_url(raw: str) -> str:
    value = (raw or "").strip()
    if not value:
        return ""
    match = _LINKTREE_TITLE_PASTE.match(value)
    if match:
        return f"https://linktr.ee/{match.group(1)}"
    value = _LINTR_TYPO.sub("https://linktr.ee/", value)
    lowered = value.lower()
    if lowered.startswith("www.linktr.ee/"):
        value = "https://" + value[4:]
    elif lowered.startswith("linktr.ee/"):
        value = "https://" + value
    return value


def read_creator_sheet(path: Path | str) -> List[Dict[str, str]]:
    """Return one dict per creator row: name, url, confirmation, language.

    confirmation/language are "" when the sheet has no such column; callers
    default confirmation to Pending (section 3).
    """
    path = Path(path)
    if path.suffix.lower() == ".csv":
        import csv

        with path.open(newline="", encoding="utf-8") as handle:
            rows = list(csv.reader(handle))
    else:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = [["" if cell is None else str(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]

    if not rows:
        return []

    header = [str(cell or "").strip().lower() for cell in rows[0]]
    name_idx = _find_column(header, NAME_COLUMNS)
    link_idx = _find_column(header, LINK_COLUMNS)
    if name_idx is None or link_idx is None:
        # No recognizable header: fall back to the first two columns.
        name_idx, link_idx = 0, 1
        confirmation_idx = language_idx = email_idx = whatsapp_idx = None
        body = rows
    else:
        confirmation_idx = _find_column(header, CONFIRMATION_COLUMNS)
        language_idx = _find_column(header, LANGUAGE_COLUMNS)
        email_idx = _find_column(header, EMAIL_COLUMNS)
        whatsapp_idx = _find_column(header, WHATSAPP_COLUMNS)
        body = rows[1:]

    def _cell(row: Sequence[str], idx: Optional[int]) -> str:
        if idx is None or idx >= len(row):
            return ""
        return str(row[idx] or "").strip()

    pairs: List[Dict[str, str]] = []
    for row in body:
        if len(row) <= max(name_idx, link_idx):
            continue
        name = _cell(row, name_idx)
        url = _clean_biolink_url(_cell(row, link_idx))
        if not name or not url:
            continue
        pairs.append(
            {
                "name": name,
                "url": url,
                "confirmation": _cell(row, confirmation_idx),
                "language": _cell(row, language_idx),
                "email": _cell(row, email_idx),
                "whatsapp": _cell(row, whatsapp_idx),
            }
        )
    return pairs


def _find_column(header: Sequence[str], options: Sequence[str]) -> Optional[int]:
    for idx, cell in enumerate(header):
        if cell in options:
            return idx
    for idx, cell in enumerate(header):
        if any(option in cell for option in options):
            return idx
    return None


_NEXT_DATA_RE = re.compile(r'<script id="__NEXT_DATA__"[^>]*>(.*?)</script>', re.S)
_LANGUAGE_KEYS = ("language", "locale", "preferredLanguage", "lang")


def _links_from_html(html: str, base_url: str) -> List[str]:
    """Pull every outbound link from a Linktree-style page.

    Linktree renders its links into a ``__NEXT_DATA__`` JSON blob as well as
    anchor tags, so read both and let the caller dedupe.
    """
    urls: List[str] = []
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        for anchor in soup.find_all("a", href=True):
            href = anchor["href"].strip()
            if href.startswith(("mailto:", "tel:", "javascript:", "#")):
                continue
            urls.append(urljoin(base_url, href))
        for script in soup.find_all("script"):
            text = script.string or ""
            if "http" in text:
                urls.extend(_URL_IN_TEXT.findall(text))
    except ImportError:  # pragma: no cover - bs4 is a declared dependency
        urls.extend(_URL_IN_TEXT.findall(html))

    # Strip JSON escaping artifacts picked up from embedded script payloads.
    cleaned = [u.rstrip("\\\",'").replace("\\u002F", "/").replace("\\/", "/") for u in urls]
    return list(dict.fromkeys(u for u in cleaned if u.startswith("http")))


def _meta_from_html(html: str) -> Dict[str, str]:
    """Avatar + language straight out of Linktree's own ``__NEXT_DATA__`` blob.

    Linktree hosts the avatar on its own CDN regardless of the creator's
    platform choices, so it's the one photo reliably available for everyone
    — same technique as Desktop/ranking's ``extract.py``.
    """
    match = _NEXT_DATA_RE.search(html)
    if not match:
        return {"avatar_url": "", "language": ""}
    try:
        page_props = json.loads(match.group(1))["props"]["pageProps"]
        account = page_props.get("account") or {}
    except (json.JSONDecodeError, KeyError, TypeError):
        return {"avatar_url": "", "language": ""}

    avatar_url = str(account.get("profilePictureUrl") or "").strip()
    language = ""
    for bucket in (account, page_props):
        for key in _LANGUAGE_KEYS:
            value = bucket.get(key) if isinstance(bucket, dict) else None
            if isinstance(value, str) and value.strip():
                language = value.strip()
                break
        if language:
            break
    return {"avatar_url": avatar_url, "language": language}


def fetch_biolink_page(url: str, fetch: Callable[[str], str]) -> Tuple[List[str], Dict[str, str]]:
    """One fetch, both the outbound links and the avatar/language metadata."""
    html = fetch(url)
    if not html:
        return [], {"avatar_url": "", "language": ""}
    return _links_from_html(html, url), _meta_from_html(html)


def scrape_biolink_page(url: str, fetch: Callable[[str], str]) -> List[str]:
    """Links only — kept for callers that don't need avatar/language too."""
    return fetch_biolink_page(url, fetch)[0]


def extract_links(
    sheet_path: Path | str,
    fetch: Callable[[str], str],
    resolver: Optional[Callable[[str], str]] = None,
) -> Tuple[List[CreatorProfile], List[Dict[str, str]]]:
    """Excel sheet -> validated creator profiles plus an unresolved log.

    A page that fails to load is logged and skipped so one dead Linktree never
    takes down the batch.
    """
    pairs = read_creator_sheet(sheet_path)
    meta_by_creator: Dict[str, Dict[str, str]] = {p["name"]: dict(p) for p in pairs}
    biolink_url_by_creator: Dict[str, str] = {}
    candidates: List[LinkCandidate] = []
    unresolved: List[Dict[str, str]] = []

    for entry in pairs:
        creator_name, biolink_url = entry["name"], entry["url"]
        biolink_url_by_creator[creator_name] = biolink_url
        try:
            found, page_meta = fetch_biolink_page(biolink_url, fetch)
        except Exception as exc:  # noqa: BLE001 - one bad page must not stop the run
            log.warning("bio-link page failed for %s (%s): %s", creator_name, biolink_url, exc)
            unresolved.append(_unresolved_row(creator_name, biolink_url, f"page fetch failed: {exc}"))
            continue
        if not found:
            unresolved.append(_unresolved_row(creator_name, biolink_url, "empty or dead bio-link page"))
            continue
        # Scraped language only fills in where the sheet didn't already say.
        meta_by_creator[creator_name]["avatar_url"] = page_meta["avatar_url"]
        if not meta_by_creator[creator_name].get("language"):
            meta_by_creator[creator_name]["language"] = page_meta["language"]
        candidates.extend(LinkCandidate(creator_name=creator_name, raw_url=u) for u in found)

    results = validate_candidates(candidates, resolver=resolver)
    profiles: List[CreatorProfile] = []
    # A creator whose page fetched fine but whose links are *all* non-social
    # (a WhatsApp button, a merch store, ...) never hits the "elif" below --
    # candidates existed, none were REJECTED, none were ACCEPTED either, so
    # nothing would get logged at all without this (found running against
    # the real roster: 10 creators out of 109 were silently invisible,
    # indistinguishable from "never processed").
    had_candidates = {c.creator_name for c in candidates}
    resolved_names: set = set()
    for result in results:
        if result.accepted:
            resolved_names.add(result.creator_name)
            meta = meta_by_creator.get(result.creator_name, {})
            profiles.append(
                CreatorProfile(
                    creator_name=result.creator_name,
                    platform=result.platform,
                    handle=result.handle,
                    profile_url=result.normalized_url,
                    status=STATUS_RESOLVED,
                    source_url=result.raw_url,
                    platform_id=result.platform_id,
                    confirmation=meta.get("confirmation") or CONFIRMATION_DEFAULT,
                    language=meta.get("language", ""),
                    avatar_url=meta.get("avatar_url", ""),
                    email=meta.get("email", ""),
                    whatsapp=meta.get("whatsapp", ""),
                )
            )
        elif result.outcome != IGNORED:
            # Non-social links are not errors (test case 7) — only log real rejects.
            unresolved.append(_unresolved_row(result.creator_name, result.raw_url, result.reason))

    # Only for creators with *zero* log entries yet -- a creator with some
    # specific rejection already logged above (e.g. an invalid Facebook
    # handle) shouldn't also get this generic fallback message.
    already_logged = {row["creator_name"] for row in unresolved}
    for creator_name in had_candidates - resolved_names - already_logged:
        unresolved.append(
            _unresolved_row(
                creator_name,
                biolink_url_by_creator.get(creator_name, ""),
                "page had links but none were recognized social platform profiles",
            )
        )

    return profiles, unresolved


def _unresolved_row(creator_name: str, url: str, reason: str) -> Dict[str, str]:
    return {
        "date": today_stamp(),
        "creator_name": creator_name,
        "url": url,
        "reason": reason,
        "stage": "extract-links",
    }


def write_profiles(repo: DataRepo, profiles: Iterable[CreatorProfile]) -> Path:
    rows = [p.to_row() for p in profiles]
    fieldnames = list(CreatorProfile("", "", "", "").to_row().keys())
    return write_csv(repo.creator_profiles, rows, fieldnames)


def revalidate_profiles(
    profiles: Iterable[CreatorProfile],
    resolver: Optional[Callable[[str], str]] = None,
) -> Tuple[List[CreatorProfile], List[Dict[str, str]]]:
    """Re-run only the URL checks over an existing creator_profiles.csv.

    Cheap enough to run whenever the source sheet changes, without re-scraping.
    """
    updated: List[CreatorProfile] = []
    unresolved: List[Dict[str, str]] = []
    for profile in profiles:
        result: ValidationResult = validate_candidates(
            [LinkCandidate(profile.creator_name, profile.profile_url)], resolver=resolver
        )[0]
        if result.accepted:
            profile.platform = result.platform
            profile.handle = result.handle
            profile.profile_url = result.normalized_url
            profile.platform_id = profile.platform_id or result.platform_id
            profile.status = STATUS_RESOLVED
            profile.reason = ""
        else:
            profile.status = STATUS_INVALID
            profile.reason = result.reason
            unresolved.append(
                {**_unresolved_row(profile.creator_name, profile.profile_url, result.reason), "stage": "validate-links"}
            )
        updated.append(profile)
    return updated, unresolved


def dump_debug(payload: object) -> str:
    """Compact JSON for the raw_snippet debug column."""
    try:
        return json.dumps(payload, ensure_ascii=False)[:500]
    except (TypeError, ValueError):
        return str(payload)[:500]
